"""Real DOCX bytes, exact fixture figures, media, ownership and CSV preservation."""
import io
import zipfile

from docx import Document
from sqlalchemy import text

from tests.test_datasets import isolated, upload, scoped, SALES, STUDENT  # noqa: F401
from tests.test_forensic_correctness import warehouse  # noqa: F401

WORD_SALES=b'date,order_id,region,revenue,status\n2025-01-01,a,Pune,200,delivered\n2025-01-02,b,Delhi,800,canceled\n'


def contents(response):
    assert response.status_code == 200, response.text[:500] if response.status_code != 200 else ''
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert 'business-sales-analysis.docx' in response.headers['content-disposition']
    assert response.content[:4] == b'PK\x03\x04'
    archive=zipfile.ZipFile(io.BytesIO(response.content))
    assert archive.testzip() is None
    assert '[Content_Types].xml' in archive.namelist()
    document=Document(io.BytesIO(response.content))
    words='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    return words,archive


def test_docx_uses_real_authorized_filter_values_and_embedded_media(isolated,analyst_auth):
    client,factory=isolated
    ds=upload(client,analyst_auth,content=WORD_SALES,name='Word verification synthetic sales')
    headers=scoped(analyst_auth,ds)
    with factory() as db:
        before=db.execute(text('SELECT COUNT(*) FROM app_dataset')).scalar()
    response=client.get('/api/reports/business-sales.docx?column_filters={"region":"Pune"}',headers=headers)
    words,archive=contents(response)
    assert 'Business Sales Analysis Report' in words
    assert 'Word verification synthetic sales' in words
    assert '200' in words and '1,000' not in words
    assert 'source revenue units' in words
    assert 'Forecast unavailable' in words and 'No replacement forecast' in words
    for section in ('Executive summary','Sales KPI table','Category and product performance','Findings and recommendations','Methodology and definitions'):
        assert section in words
    assert len([n for n in archive.namelist() if n.startswith('word/media/')])>=3
    assert 'Selected observed coverage: 2025-01-01 to 2025-01-01' in words
    assert ds['dataset_id'] not in words
    assert 'C:\\Users\\' not in words
    with factory() as db:
        assert db.execute(text('SELECT COUNT(*) FROM app_dataset')).scalar()==before
    csv=client.get('/api/reports/export?report=dataset_rows&column_filters={"region":"Pune"}',headers=headers)
    assert csv.status_code==200 and 'Pune' in csv.text and 'Delhi' not in csv.text


def test_word_requires_auth_and_preserves_dataset_ownership(isolated,analyst_auth,auth):
    client,_=isolated
    ds=upload(client,auth,name='Private admin source')
    assert client.get('/api/reports/business-sales.docx',headers=scoped(analyst_auth,ds)).status_code==403
    assert client.get('/api/reports/business-sales.docx',headers={'X-Dataset-ID':ds['dataset_id']}).status_code==401
    assert client.get('/api/reports/business-sales.docx?date_from=bad',headers=scoped(auth,ds)).status_code==422


def test_generic_dataset_gets_honest_profile_not_fake_sales(isolated,analyst_auth):
    client,_=isolated
    ds=upload(client,analyst_auth,content=STUDENT,name='Generic educational profile')
    words,archive=contents(client.get('/api/reports/business-sales.docx',headers=scoped(analyst_auth,ds)))
    assert 'generic dataset profiling report' in words
    assert 'not invented' in words
    assert len([n for n in archive.namelist() if n.startswith('word/media/')])==1


def test_bundle_report_keeps_distinct_order_chart_despite_member_row_count(isolated,analyst_auth):
    from tests.test_olist_bundle import upload as upload_bundle
    client,_=isolated
    response=upload_bundle(client,analyst_auth)
    assert response.status_code==201
    ds=response.json()
    assert ds['row_count']==21
    words,_=contents(client.get('/api/reports/business-sales.docx',headers=scoped(analyst_auth,ds)))
    assert 'Observed monthly distinct orders using the same report scope.' in words
    assert 'Counts of distinct matching orders by current recorded status.' in words


def test_failed_report_releases_capacity(monkeypatch):
    from app.services import business_report
    def fail(*args): raise RuntimeError('fixture render failure')
    monkeypatch.setattr(business_report,'snapshot',fail)
    import pytest
    for _ in range(3):
        with pytest.raises(RuntimeError,match='fixture render failure'):business_report.generate(None,None)


def test_report_execution_hint_preserves_existing_evidence_exactly(warehouse):
    from app.analytics.business import evidence
    from app.analytics.filters import Filters
    from types import SimpleNamespace
    warehouse.info['dataset']=SimpleNamespace(id='reference',source_type='warehouse',adapter='star_schema',semantics={'currency_code':'BRL'})
    for domain in ('customers','sellers','revenue'):
        assert evidence(warehouse,Filters(),domain)==evidence(warehouse,Filters(),domain,materialize=True)


def test_duplicate_order_rows_do_not_become_false_cancellation_zeroes(isolated,analyst_auth):
    client,_=isolated
    raw=b'date,order_id,revenue,status\n2025-01-01,a,10,delivered\n2025-01-01,a,20,delivered\n'
    ds=upload(client,analyst_auth,content=raw,name='Item-like rows')
    words,_=contents(client.get('/api/reports/business-sales.docx',headers=scoped(analyst_auth,ds)))
    assert 'Cancellation rate is unavailable' in words
    assert '30' in words


def test_sql_deadline_has_an_actionable_error_and_releases_capacity(monkeypatch):
    from app.services import business_report
    from types import SimpleNamespace
    from sqlalchemy.exc import OperationalError
    from fastapi import HTTPException
    import pytest
    def interrupted(*args):raise OperationalError('interrupted',{},Exception())
    monkeypatch.setattr(business_report,'snapshot',interrupted)
    for _ in range(3):
        ticks=iter([0,121]);monkeypatch.setattr(business_report,'time',SimpleNamespace(monotonic=lambda:next(ticks)))
        with pytest.raises(HTTPException) as error:business_report.generate(None,None)
        assert error.value.status_code==503 and 'narrower date range' in error.value.detail
