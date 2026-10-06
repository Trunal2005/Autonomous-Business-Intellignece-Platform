"""Bounded ingestion, schema intelligence, registry and authorized resolution.

Uploaded datasets live in isolated typed tables in the existing database.
Only server-generated identifiers enter SQL; source names remain metadata.
"""
from __future__ import annotations

import csv
import io
import json
import math
import os
import re
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from fastapi import HTTPException
from sqlalchemy import Column, Float, Integer, MetaData, String, Table, Text, func, inspect, select
from sqlalchemy.orm import Session

from app.models.dataset import Dataset, DatasetSelection

MAX_BYTES = int(os.getenv("DATASET_MAX_BYTES", 20 * 1024 * 1024))
MAX_ROWS = int(os.getenv("DATASET_MAX_ROWS", 100_000))
MAX_COLUMNS = int(os.getenv("DATASET_MAX_COLUMNS", 100))
MAX_SECONDS = int(os.getenv("DATASET_MAX_SECONDS", 60))
MAX_CELLS = int(os.getenv("DATASET_MAX_CELLS", 2_000_000))
REFERENCE_ID = "00000000-0000-0000-0000-000000000001"

# Deliberately conservative: ambiguous 'amount' is currency, not revenue;
# arbitrary identifiers and numeric fields never become customers or sales.
ALIASES = {
    "revenue": {"revenue", "sales", "salesamount", "totalrevenue", "itemrevenue"},
    "currency": {"amount", "totalamount", "transactionamount", "paymentvalue"},
    "price": {"price", "unitprice", "sellingprice"},
    "quantity": {"quantity", "qty", "units", "unitssold", "nitems"},
    "customer": {"customer", "customerid", "customeruniqueid", "custid", "clientid"},
    "user": {"userid", "accountid"},
    "product": {"product", "productid", "item", "itemid", "sku", "service"},
    "seller": {"seller", "sellerid", "vendor", "vendorid", "merchantid"},
    "order": {"orderid", "order", "transactionid", "invoiceid"},
    "category": {"category", "productcategory", "productcategoryname", "department"},
    "location": {"region", "location", "city", "state", "country", "customerstate"},
    "seller_location": {"sellerstate", "vendorstate"},
    "status": {"status", "orderstatus", "transactionstatus"},
    "payment_type": {"paymenttype", "paymentmethod"},
    "rating": {"rating", "reviewscore", "rating score"},
    "delivery_date": {"deliverydate", "delivereddate", "deliveredcustomerdate"},
    "estimated_date": {"estimateddeliverydate", "estimateddate"},
    "weight_g": {"productweightg", "weightg"},
}
ROLES = set(ALIASES) | {"date", "identifier", "numeric_measure", "percentage", "text", "boolean"}
REFERENCE_FIELDS = {"date": "purchase_date", "order": "order_id", "customer": "customer_unique_id",
                    "revenue": "item_revenue", "status": "order_status", "rating": "review_score",
                    "delivery_date": "delivered_customer_date", "estimated_date": "estimated_delivery_date"}


def reference_profile(db):
    columns = inspect(db.get_bind()).get_columns("fact_orders")
    return {"grain": "orders", "columns": [
        {"name": c["name"], "dtype": str(c["type"]), "role": next((r for r, n in REFERENCE_FIELDS.items() if n == c["name"]), "reference field"),
         "confidence": 1, "missing": None, "unique": None, "constant": False, "ambiguous_date": False} for c in columns]}


def key(name):
    return re.sub(r"[^a-z0-9]", "", re.sub(r"([a-z])([A-Z])", r"\1_\2", str(name)).lower())


def ensure_registry(db: Session):
    bind = db.get_bind()
    Dataset.__table__.create(bind, checkfirst=True)
    DatasetSelection.__table__.create(bind, checkfirst=True)
    if not db.get(Dataset, REFERENCE_ID) and "fact_orders" in inspect(bind).get_table_names():
        from sqlalchemy import text
        n = db.scalar(text("SELECT COUNT(*) FROM fact_orders"))
        if n:
            profile = reference_profile(db)
            db.add(Dataset(id=REFERENCE_ID, name="Olist", owner=None, status="READY",
                           source_type="warehouse", adapter="star_schema", row_count=n,
                           column_count=len(profile["columns"]), profile=profile,
                           semantics={"currency_code": "BRL", "domain": "commerce", "fields": REFERENCE_FIELDS},
                           capabilities={s: {"available": True, "reason": None} for s in
                                         ("dashboard", "sales", "orders", "customers", "products", "sellers", "delivery")}))
            db.commit()
    reference = db.get(Dataset, REFERENCE_ID)
    if reference and not reference.profile.get("columns"):
        reference.profile = reference_profile(db)
        reference.column_count = len(reference.profile["columns"])
        reference.semantics = {**reference.semantics, "fields": REFERENCE_FIELDS}
        db.commit()


def public(ds):
    return {"dataset_id": ds.id, "dataset_name": ds.name, "owner": ds.owner,
            "status": ds.status, "source_type": ds.source_type,
            "created_at": ds.created_at.isoformat() if ds.created_at else None,
            "updated_at": ds.updated_at.isoformat() if ds.updated_at else None,
            "row_count": ds.row_count, "column_count": ds.column_count,
            "schema": {**ds.profile, "columns": [
                {k: v for k, v in column.items() if k != "storage"}
                for column in ds.profile.get("columns", [])]}, "semantics": ds.semantics,
            "capabilities": ds.capabilities, "quality": ds.quality, "error": ds.error,
            "read_only": ds.owner is None}


def authorized(db, dataset_id, user, ready=False):
    ds = db.get(Dataset, dataset_id)
    if ds is None:
        raise HTTPException(404, "Dataset not found.")
    # Admin privileges do not implicitly grant access to private analyst data.
    if ds.owner is not None and ds.owner != user["username"]:
        raise HTTPException(403, "You do not have access to this dataset.")
    if ready and ds.status != "READY":
        raise HTTPException(409, "Dataset is not ready. Review its processing status.")
    return ds


def resolve(db, user, requested=None):
    ensure_registry(db)
    selection = db.get(DatasetSelection, user["username"])
    dataset_id = requested or (selection.dataset_id if selection else None)
    if not dataset_id:
        # Migration initializes existing users to the registered reference once.
        # A removed or failed selection is never replaced by reference data.
        if db.get(Dataset, REFERENCE_ID):
            selection = DatasetSelection(username=user["username"], dataset_id=REFERENCE_ID)
            db.add(selection)
            db.commit()
            dataset_id = REFERENCE_ID
        else:
            raise HTTPException(409, "Select a ready dataset in Dataset Manager.")
    ds = authorized(db, dataset_id, user, ready=True)
    db.info["dataset"] = ds
    return ds


def storage_table(ds, db=None, columns=None):
    # Revalidate even persisted IDs before constructing database identifiers.
    ident = uuid.UUID(ds.id).hex
    name = "dataset_" + ident
    if columns is None:
        tables = db.info.setdefault("dataset_tables", {})
        if ds.id not in tables:
            tables[ds.id] = Table(name, MetaData(), autoload_with=db.connection())
        return tables[ds.id]
    return Table(name, MetaData(), Column("row_id", Integer, primary_key=True), *[
        Column(f"c{i}", Float if c["dtype"] == "numeric" else Text, nullable=True)
        for i, c in enumerate(columns)
    ])


def _headers(values):
    names = [str(v).strip() if v is not None else "" for v in values]
    if not names or any(not n or n.lower().startswith("unnamed:") for n in names):
        raise ValueError("Missing column headers were detected.")
    if len(names) > MAX_COLUMNS:
        raise ValueError(f"Dataset exceeds the {MAX_COLUMNS} column limit.")
    if len({n.casefold() for n in names}) != len(names):
        raise ValueError("Duplicate column names were detected.")
    if any(len(n) > 200 for n in names):
        raise ValueError("Column names must be at most 200 characters.")
    if all(re.fullmatch(r"[\d.]+", n) for n in names):
        raise ValueError("A header row is required.")
    return names


def parse_source(raw, suffix):
    if not raw:
        raise ValueError("The uploaded file is empty.")
    if len(raw) > MAX_BYTES:
        raise ValueError(f"File exceeds the {MAX_BYTES // 1024 // 1024} MB limit.")
    if suffix == ".csv":
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV must use UTF-8 encoding. Please export it as UTF-8.") from None
        if "\x00" in content:
            raise ValueError("CSV contains invalid binary content.")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(content), dialect, strict=True)
        names = _headers(next(reader, []))
        rows = []
        for row in reader:
            if not row:
                continue
            if len(row) != len(names):
                raise ValueError("CSV rows must have the same number of fields as the header.")
            if len(rows) >= MAX_ROWS:
                raise ValueError(f"Dataset exceeds the {MAX_ROWS} row limit.")
            if (len(rows) + 1) * len(names) > MAX_CELLS:
                raise ValueError("Dataset exceeds the configured processing memory budget.")
            rows.append(row)
        df = pd.DataFrame(rows, columns=names)
    elif suffix == ".json":
        def pairs(items):
            result = {}
            for k, v in items:
                if k in result:
                    raise ValueError("Duplicate JSON field names were detected.")
                result[k] = v
            return result
        obj = json.loads(raw.decode("utf-8-sig"), object_pairs_hook=pairs,
                         parse_constant=lambda _: (_ for _ in ()).throw(ValueError("Non-finite JSON numbers are unsupported.")))
        if not isinstance(obj, list) or not obj or not all(isinstance(r, dict) for r in obj):
            raise ValueError("JSON must contain a non-empty array of flat records.")
        if len(obj) > MAX_ROWS:
            raise ValueError(f"Dataset exceeds the {MAX_ROWS} row limit.")
        if any(isinstance(v, (dict, list)) for r in obj for v in r.values()):
            raise ValueError("Nested JSON fields are unsupported; use flat records.")
        names = _headers(list(dict.fromkeys(k for r in obj for k in r)))
        if len(names) * len(obj) > MAX_CELLS:
            raise ValueError("Dataset exceeds the configured processing memory budget.")
        df = pd.DataFrame(obj)
        df.columns = _headers(df.columns)
    elif suffix in {".xlsx", ".xls"}:
        if suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_BYTES * 10:
                    raise ValueError("Spreadsheet expanded size exceeds the safety limit.")
                if any("vbaProject" in i.filename or "externalLinks/" in i.filename for i in archive.infolist()):
                    raise ValueError("Macros and external spreadsheet links are unsupported.")
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                if len(workbook.worksheets) != 1:
                    raise ValueError("Upload a workbook with one worksheet.")
                sheet = workbook.worksheets[0]
                if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) * (sheet.max_row or 0) > MAX_CELLS:
                    raise ValueError("Spreadsheet dimensions exceed the configured processing limits.")
            finally:
                workbook.close()
        try:
            book = pd.ExcelFile(io.BytesIO(raw), engine="openpyxl" if suffix == ".xlsx" else "xlrd")
        except ImportError:
            raise ValueError("Legacy XLS support requires xlrd; export this workbook as XLSX or CSV.") from None
        if len(book.sheet_names) != 1:
            raise ValueError("Upload a workbook with one worksheet. Separate sheets are separate datasets.")
        frame = pd.read_excel(book, header=None, nrows=MAX_ROWS + 2)
        names = _headers(frame.iloc[0].tolist() if not frame.empty else [])
        df = frame.iloc[1:].copy()
        df.columns = names
    else:
        raise ValueError("Unsupported file type. Use CSV, XLSX, XLS or flat-record JSON.")
    if df.empty:
        raise ValueError("Dataset contains no data rows.")
    if len(df) > MAX_ROWS or len(df.columns) > MAX_COLUMNS or len(df) * len(df.columns) > MAX_CELLS:
        raise ValueError("Dataset exceeds the configured row or column limit.")
    return df.replace(r"^\s*$", None, regex=True).replace({"NaN": None, "null": None})


def profile_frame(df, date_order=None):
    cols, warnings = [], []
    normalized = df.copy()
    for i, name in enumerate(df.columns):
        s = df[name]
        present = s.dropna()
        canonical = key(name)
        currency_code = next((code for code in ("BRL", "USD", "INR", "EUR", "GBP")
                              if canonical.endswith(code.lower()) and canonical[:-3] in ALIASES["revenue"] | ALIASES["currency"] | ALIASES["price"]), None)
        semantic_key = canonical[:-3] if currency_code else canonical
        idlike = canonical.endswith("id") or canonical in {"phone", "phonenumber", "zipcode", "postalcode"}
        role = next((r for r, aliases in ALIASES.items() if semantic_key in aliases), None)
        if canonical in {"currency", "currencycode"}:
            codes = set(present.astype(str).str.upper().str.strip())
            currency_code = next(iter(codes)) if len(codes) == 1 and codes <= {"BRL", "USD", "INR", "EUR", "GBP"} else None
            role = "text"
        numeric = pd.to_numeric(s, errors="coerce")
        boolean = bool(len(present) and set(present.astype(str).str.lower()) <= {"true", "false", "yes", "no"})
        is_num = len(present) > 0 and numeric.notna().sum() == len(present) and not idlike and not boolean
        dtype, confidence = "text", 0.6
        ambiguous = False
        date_candidate = any(t in canonical for t in ("date", "timestamp", "datetime"))
        date_candidate |= bool(len(present) and present.astype(str).str.match(r"^\d{4}-\d{2}-\d{2}(?:[ T].*)?$").all())
        if date_candidate and len(present):
            slash = present.astype(str).str.extract(r"^(\d{1,2})[/-](\d{1,2})[/-]\d{4}$").dropna()
            ambiguous = bool(not slash.empty and (slash.astype(int) <= 12).all().all() and not date_order)
            if ambiguous:
                warnings.append(f"Choose day/month order for {name} before using it as an analytical date.")
            else:
                first_large = not slash.empty and (slash[0].astype(int) > 12).any()
                second_large = not slash.empty and (slash[1].astype(int) > 12).any()
                if first_large and second_large and not date_order:
                    dates = pd.Series(pd.NaT, index=s.index)
                else:
                    dates = pd.to_datetime(s.astype("string"), errors="coerce", format="mixed",
                                           dayfirst=date_order == "dayfirst" or (date_order is None and first_large), utc=True)
                if dates.notna().sum() == len(present):
                    dtype, confidence = "datetime", 0.98
                    role = role if role in {"delivery_date", "estimated_date"} else "date"
                    normalized[name] = dates.map(lambda d: d.isoformat() if pd.notna(d) else None)
                else:
                    warnings.append(f"{name} contains invalid date values; date analytics are disabled for this field.")
        if dtype == "text" and is_num and not date_candidate:
            if not all(math.isfinite(float(v)) for v in numeric.dropna()):
                raise ValueError(f"{name} contains non-finite numeric values.")
            dtype, confidence = "numeric", 0.9
            normalized[name] = numeric
        if dtype == "text" and not role:
            if idlike:
                role = "identifier"
            elif boolean:
                role, dtype = "boolean", "boolean"
            elif s.nunique() <= min(50, max(2, len(s) // 2)):
                role = "category"
            else:
                role = "text"
        if dtype == "numeric" and not role:
            role = "percentage" if any(t in canonical for t in ("percent", "percentage", "attendance")) else "numeric_measure"
        # Monetary/quantity roles require numeric evidence; timestamps require parsed values.
        if role in {"revenue", "currency", "price", "quantity", "rating", "weight_g"} and dtype != "numeric":
            role = "text"
        if role in {"delivery_date", "estimated_date", "date"} and dtype != "datetime":
            role = "text"
        c = {"name": str(name), "storage": f"c{i}", "dtype": dtype, "role": role,
             "currency_code": currency_code,
             "confidence": round(confidence if role in {"text", "category", "numeric_measure"} else min(confidence, .95), 2),
             "missing": int(s.isna().sum()), "missing_percentage": round(float(s.isna().mean() * 100), 2),
             "unique": int(s.nunique()), "constant": bool(s.nunique() <= 1), "ambiguous_date": ambiguous}
        if date_candidate:
            c["invalid_dates"] = int(len(present) - dates.notna().sum()) if not ambiguous and len(present) else 0
        if dtype == "numeric":
            c["statistics"] = {k: (float(v) if pd.notna(v) else None) for k, v in
                               {"min": numeric.min(), "max": numeric.max(), "mean": numeric.mean(),
                                "median": numeric.median(), "std": numeric.std()}.items()}
        cols.append(c)
    return normalized, cols, {"missing_values": int(df.isna().sum().sum()),
                              "duplicate_rows": int(df.duplicated().sum()), "warnings": warnings}


def infer_semantics(cols):
    mappings, ambiguous = {}, []
    for role in ROLES - {"identifier", "text", "boolean", "numeric_measure", "percentage"}:
        candidates = [c for c in cols if c["role"] == role]
        if len(candidates) == 1:
            mappings[role] = candidates[0]["name"]
        elif candidates:
            ambiguous.append({"role": role, "candidates": [c["name"] for c in candidates]})
    codes = {c.get("currency_code") for c in cols if c.get("currency_code")}
    return {"fields": mappings, "ambiguities": ambiguous, "currency_code": next(iter(codes)) if len(codes) == 1 else None,
            "domain": "commerce" if "revenue" in mappings or "product" in mappings else "general"}


def detect_capabilities(semantics):
    fields = dict(semantics.get("fields", {}))
    if "user" in fields:
        fields.setdefault("customer", fields["user"])
    requirements = {"dashboard": [], "sales": [], "orders": ["order"], "customers": ["customer"],
                    "products": ["product"], "sellers": ["seller"], "delivery": ["date", "delivery_date"]}
    return {name: {"available": all(f in fields for f in needed),
                   "reason": None if all(f in fields for f in needed) else
                   "No " + ", ".join(f.replace("_", " ") for f in needed if f not in fields) + " field was confirmed."}
            for name, needed in requirements.items()}


def ingest(db, user, name, raw, suffix):
    ensure_registry(db)
    ds = Dataset(id=str(uuid.uuid4()), name=name, owner=user["username"], status="PROCESSING",
                 source_type=suffix.lstrip("."), raw_source=raw)
    db.add(ds)
    db.commit()
    started = time.monotonic()
    table = None
    try:
        source = parse_source(raw, suffix)
        normalized, columns, quality = profile_frame(source)
        if time.monotonic() - started > MAX_SECONDS:
            raise ValueError("Dataset processing exceeded the time limit.")
        table = storage_table(ds, columns=columns)
        table.create(db.connection())
        # Small insert batches keep memory and statement size bounded.
        for start in range(0, len(normalized), 1000):
            rows = []
            for values in normalized.iloc[start:start + 1000].itertuples(index=False, name=None):
                rows.append({f"c{i}": None if pd.isna(v) else float(v) if columns[i]["dtype"] == "numeric" else str(v)
                             for i, v in enumerate(values)})
            db.execute(table.insert(), rows)
            if time.monotonic() - started > MAX_SECONDS:
                raise ValueError("Dataset processing exceeded the time limit.")
        for col in columns:
            if col["role"] in {"date", "category", "customer", "product", "seller", "order", "location", "status"}:
                from sqlalchemy import Index
                Index(f"ix_{table.name}_{col['storage']}", table.c[col["storage"]]).create(db.connection())
        ds.row_count, ds.column_count = len(source), len(columns)
        ds.profile, ds.quality = {"columns": columns}, {**quality, "processing_seconds": round(time.monotonic() - started, 3)}
        ds.semantics = infer_semantics(columns)
        ds.capabilities = detect_capabilities(ds.semantics)
        ds.status = "READY"
        db.commit()
    except Exception as exc:
        db.rollback()
        if table is not None:
            table.drop(db.connection(), checkfirst=True)
        ds = db.get(Dataset, ds.id)
        ds.status = "FAILED"
        ds.error = str(exc)[:400] if isinstance(exc, (ValueError, csv.Error, json.JSONDecodeError)) else "Dataset could not be processed. Check the file format and structure."
        db.commit()
    return ds
