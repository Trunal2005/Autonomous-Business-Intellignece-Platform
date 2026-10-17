"""Pure DOCX builder from a bounded numeric snapshot and generated PNG media."""
import io
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))[:2000]


def number(value):
    return 'Unavailable' if value is None else f'{value:,.2f}'.rstrip('0').rstrip('.') if isinstance(value,(int,float)) else clean(value)


def build(data, media):
    doc=Document(); section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.65);section.left_margin=section.right_margin=Inches(.75)
    for style in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        s=doc.styles[style];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
        s.font.size=Pt(11 if style=='Normal' else 24 if style=='Title' else 15 if style=='Heading 1' else 12)
        s.paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    section.header.paragraphs[0].text='ABI  |  Business sales analysis'
    footer=section.footer.paragraphs[0];footer.text='Confidential aggregate analysis  |  Page '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.author='Autonomous Business Intelligence Platform'
    doc.core_properties.title='Business Sales Analysis Report'
    def paragraph(value,style=None): return doc.add_paragraph(clean(value),style)
    def heading(value): doc.add_heading(value,level=1)
    def table(headers,rows):
        if not rows:
            paragraph('No applicable records are available for this table.');return
        t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        widths=[7/len(headers)]*len(headers)
        if len(headers)==2:widths=[4.5,2.5]
        if len(headers)==3:widths=[3.1,1.7,2.2]
        for i,text in enumerate(headers):t.rows[0].cells[i].text=clean(text)
        for row in rows[:20]:
            cells=t.add_row().cells
            for i,value in enumerate(row):cells[i].text=number(value)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for r,row in enumerate(t.rows):
            for i,cell in enumerate(row.cells):
                cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    element=OxmlElement('w:'+edge);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');borders.append(element)
                props.append(borders)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF' if r==0 else 'F4F6F8' if r%2==0 else 'FFFFFF');props.append(shade)
                margins=OxmlElement('w:tcMar')
                for side in ('top','left','bottom','right'):
                    margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'90');margin.set(qn('w:type'),'dxa');margins.append(margin)
                props.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.keep_with_next=r<2
                    for run in p.runs:run.font.size=Pt(10);run.bold=r==0
        paragraph('')
    def image(key,caption):
        if key in media:
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(io.BytesIO(media[key]),width=Inches(6.8))
            p=paragraph(caption);p.paragraph_format.space_after=Pt(10)
            for run in p.runs:run.italic=True;run.font.size=Pt(9)
        else:paragraph('Chart unavailable because the selected data does not support this metric.')
    k=data['kpis']; unit=data['money_unit']; scope=data['scope']
    filters={key:('Custom column filter applied' if key=='column_filters' else value) for key,value in data['filters'].items()}
    summary='; '.join(f'{key.replace("_"," ")}: {value}' for key,value in filters.items()) or 'Full selected dataset history; no additional filters'
    paragraph('Business Sales Analysis Report','Title');paragraph(data['dataset'],'Subtitle')
    paragraph(scope);paragraph('Generated '+data['generated_at'])
    paragraph('Applied scope  '+summary)
    paragraph('Synthetic extended data' if data['synthetic'] else 'Source provenance unconfirmed','Heading 2')
    paragraph(data['provenance'])
    paragraph('This report presents the selected data’s calculated sales results, observed patterns and investigation recommendations. Model estimates are separated from realized sales. Missing or unsupported information remains unavailable.')
    paragraph('Confidential aggregate analysis. Customer identifiers, private source paths and database identifiers are excluded.')
    doc.add_page_break()
    heading('Executive summary')
    if data['sales_supported']:
        paragraph(f"The selected scope contains {number(k.get('total_orders'))} distinct orders and {number(k.get('total_revenue'))} item revenue in {unit}. Average order value is {number(k.get('avg_order_value'))}. These figures describe the recorded selected transactions and do not establish profit or causal business effects.")
        paragraph(f"Observed customers: {number(k.get('unique_customers'))}. Item units: {number(k.get('items_sold'))}. Active sellers: {number(k.get('active_sellers'))}. Definitions and eligibility limits appear below.")
    else:paragraph('This selection does not provide compatible sales semantics. This is a generic dataset profiling report; sales, customer and forecasting figures are not invented.')
    if data['recommendations']:paragraph('Priority investigation  '+data['recommendations'][0]['observation']+' '+data['recommendations'][0]['suggested_action'])
    else:paragraph('No supported actionable recommendation was generated for this scope. Review the data coverage and unavailable fields before drawing conclusions.')
    heading('Dataset scope and data quality')
    paragraph(f"{scope}. Imported source rows: {number(data['source_rows'])}. Tables: {number(data['table_count'])}. Selected observed coverage: {data['coverage'].get('min') or 'unavailable'} to {data['coverage'].get('max') or 'unavailable'}. {summary}.")
    paragraph('Full source coverage: '+str(data.get('source_coverage',{}).get('min','unavailable'))+' to '+str(data.get('source_coverage',{}).get('max','unavailable'))+'. Selected-period totals and charts use only the applied filters.')
    paragraph(data['provenance']+' Monetary unit: '+unit+'. Zero, missing and not applicable are distinct. Unknown dates and missing prices are not automatically zero activity. Filtered results may exclude valid records outside the selected scope.')
    if data.get('coverage_stats'):
        counts=data['coverage_stats']
        paragraph(f"Selected scope contains {counts['observed_days']} observed daily dates, {number(counts['missing_calendar_days'])} missing calendar dates and {counts['unknown_revenue_days']} observed dates with unknown revenue. Zero-filled dates: 0. Observed dates do not independently certify transaction completeness."+(' The owner declared this extended source coverage unknown.' if data['synthetic'] else ' No source coverage confirmation is assumed.'))
    if data.get('quality'):paragraph('Measured source quality notes  '+clean(data['quality']))
    heading('Sales KPI table')
    definitions={'total_orders':'Distinct matching orders','total_revenue':'Qualifying item prices excluding freight and payments','avg_order_value':'Qualifying item revenue / matching distinct orders','items_sold':'Observed item units','unique_customers':'Distinct validated stable customer identity','active_sellers':'Distinct sellers on matching items','avg_delivery_days':'Mean valid observed completion duration in days'}
    rows=[[key.replace('_',' ').title(),k.get(key),definition] for key,definition in definitions.items()] if data['sales_supported'] else [[p['label'],p.get('value'),'Selected table profile'] for p in data['metric_cards'][:12]]
    table(['Metric','Value','Definition'],rows)
    paragraph('Revenue and AOV use '+unit+'. Profit, margin, ROI and inventory performance are unavailable without validated cost and stock observations.')
    heading('Revenue and orders over time')
    image('revenue','Observed monthly item revenue using the report filters. Unit: '+unit+'. No forecast values enter these totals.')
    image('orders','Observed monthly distinct orders using the same report scope. Missing periods are not fabricated.')
    revenue=data['evidence'].get('revenue',{})
    cards={p['label']:p['value'] for p in revenue.get('cards',[])}
    previous,recent=cards.get('Previous 14-day revenue'),cards.get('Recent 14-day revenue')
    if previous is not None and recent is not None:
        paragraph(f"Comparable observed windows: {revenue.get('measurement_window')}. Revenue was {number(previous)} then {number(recent)} {unit}."+(f' The measured change is {(recent/previous-1)*100:.2f}%.' if previous else ' Relative change is undefined because the preceding total is zero.'))
    else:paragraph('No comparable complete 14-day-window trend is available. Partial months are not treated as equal exposure periods.')
    heading('Category and product performance')
    image('categories','Top ten categories by qualifying item revenue. Full selected-scope revenue is the contribution denominator.')
    total=k.get('total_revenue')
    table(['Category','Revenue','Share of scope'],[[p.get('category') or 'Unmapped',p.get('revenue'),f"{100*p['revenue']/total:.2f}%" if total and p.get('revenue') is not None else 'Unavailable'] for p in data['categories']])
    for key,title in [('by_revenue','Leading products by revenue'),('by_quantity','Leading products by item units')]:
        products=data.get('products',{}).get(key,[])
        if products:
            paragraph(title,'Heading 2');table(['Aggregate product group','Revenue','Item units'],[[f'Product group {i+1}',p.get('revenue'),p.get('quantity')] for i,p in enumerate(products)])
    heading('Order status and cancellations')
    image('statuses','Counts of distinct matching orders by current recorded status.')
    orders=k.get('total_orders')
    table(['Status','Orders','Share of orders'],[[p.get('status') or 'Unknown',p.get('orders'),f"{100*p['orders']/orders:.2f}%" if orders else 'Unavailable'] for p in data['statuses']])
    canceled=sum(p.get('orders',0) for p in data['statuses'] if str(p.get('status') or '').casefold() in ('canceled','cancelled'))
    paragraph(f"Cancellation rate: {100*canceled/orders:.2f}% ({canceled:,} canceled / {orders:,} matching orders)." if orders and data['statuses'] else 'Cancellation rate is unavailable without validated order-grain status counts and a matching-order denominator.')
    heading('Customer analysis')
    customer=data['evidence'].get('customers',{})
    table(['Observed customer metric','Value'],[[p['label'],p['value']] for p in customer.get('cards',[])])
    image('customers','Purchase frequency is based on distinct matching orders per stable customer; it is historical, not a future customer forecast.')
    image('customer_revenue','First-time versus returning revenue uses the first observed purchase in complete available history. It does not establish lifetime acquisition or forecast retention.')
    for note in customer.get('limitations',[]):paragraph(note)
    heading('Seller and delivery performance')
    seller=data['evidence'].get('sellers',{})
    table(['Eligible delivery metric','Value'],[[p['label'],p['value']] for p in seller.get('cards',[])])
    image('delivery','Rates use eligible delivered orders with valid purchase, actual delivery and estimate dates. Missing timestamps are not zero-day deliveries.')
    image('days','Mean valid completed delivery duration by matching purchase month, in days.')
    image('cancellation_rate','Canceled matching orders divided by all matching orders, by purchase month. Unit: percent.')
    table(['Seller aggregate group','Eligible orders','Late rate'],[[f'Seller group {i+1}',p.get('orders'),f"{p['value']:.2f}%"] for i,p in enumerate(seller.get('rankings',[])[:10])])
    for note in seller.get('limitations',[]):paragraph(note)
    heading('Forecast and machine learning findings')
    prediction=data.get('forecast')
    if prediction:
        model=prediction.get('model',{});future=prediction['forecast']
        paragraph(f"Model: {model.get('algorithm')} version {model.get('version')}. Origin: {prediction.get('forecast_origin')}. Forecast: {future[0]['date']} to {future[-1]['date']}, {len(future)} daily expected-order estimates. These are model outputs, not realized sales.")
        image('forecast','Solid observed order history and dashed future model estimates. No calibrated uncertainty interval is available.')
        for warning in prediction.get('warnings',[]):paragraph(warning)
        evaluation=data.get('forecast_evaluation')
        if evaluation and evaluation.get('evaluation'):
            paragraph('Offline synthetic source evaluation is separate from live filtered-scope accuracy. Validation uses 2020; reserved holdout uses 2021, before the final refit. Real-world performance is unverified.')
            table(['Holdout horizon','Last value MAE','Weekly naive MAE','Selected model MAE'],[[h+' days',v['metrics']['last_value']['mae'],v['metrics']['seasonal_naive']['mae'],v['metrics'][evaluation['algorithm']]['mae']] for h,v in evaluation['evaluation']['holdout'].items()])
            for reason in evaluation['quality']['reasons']:paragraph('Qualification limitation  '+reason)
    else:paragraph('Forecast unavailable  '+str(data.get('forecast_reason'))+'. No replacement forecast or interval was fabricated.')
    heading('Findings and recommendations')
    if not data['recommendations']:paragraph('No actionable signal is supported under the current deterministic evidence rules.')
    for item in data['recommendations']:
        paragraph(item['title'],'Heading 2');paragraph(item['observation']);paragraph('Measurement window  '+item['measurement_window']);paragraph('Suggested action  '+item['suggested_action']);paragraph('Limitation  '+item['limitations'])
    heading('Methodology and definitions')
    paragraph('All numeric inputs are authorized selected-data aggregates. Order-item amounts are aggregated at item grain without multiplying them through payment/review joins. Revenue excludes freight and is not added to payment totals. AOV uses the same matching-order denominator as the revenue scope. Categories are top ten with the complete-scope denominator. Product tables show up to five leaders per metric; seller rankings show up to ten groups. Generic numeric profiling shows up to twelve fields. Detail rows are never sampled to calculate aggregate denominators.')
    paragraph('Customer repeat rate uses at least two matching distinct orders per validated customer identity. First observed purchase uses all available selected-dataset history and is not lifetime acquisition. Seller rankings require at least twenty eligible distinct delivered orders. Late means actual delivery after the estimated timestamp; association does not establish seller causation. Cancellations divide canceled matching orders by all matching orders.')
    paragraph('Synthetic forecasts use only registered dataset bindings and past-only features. Unknown daily observations stay missing. Their quality is conditional on the synthetic source and cannot validate real-market performance. Legacy forecast fit provenance remains unverified. Earlier-origin comparisons with a model refitted through the source end are not independent backtests.')
    image('flow','Actual report workflow. Training is an explicit offline action and is never triggered by report generation.')
    heading('Appendix')
    paragraph('Scope summary  '+summary+'. Source unit  '+unit+'. Generation time  '+data['generated_at']+'. No raw customer records, private filesystem locations or internal dataset identifiers are included.')
    for row in data.get('numeric_statistics',[])[:12]:
        numeric={k:v for k,v in row.items() if k in ('column','field','count','min','max','mean','missing')}
        if numeric:paragraph('Available numeric profile  '+str(numeric))
    for border in list(doc.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    output=io.BytesIO();doc.save(output);return output.getvalue()
